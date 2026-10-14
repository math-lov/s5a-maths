# -*- coding: utf-8 -*-
"""由網站資料檔 data/src/<part>.json 生成 Word 工作紙（實體列印用）。

用法：
    python -X utf8 tools/make_worksheet_docx.py                      # 預設 ch17-test
    python -X utf8 tools/make_worksheet_docx.py --part ch17-qc
    python -X utf8 tools/make_worksheet_docx.py --part ch17-test --out docs/worksheets/xxx.docx

為甚麼由資料檔生成：工作紙與網站題目共用同一份 data/src/<part>.json，
所以改題目（字句、分數、題數）只要改 JSON 再跑一次，兩邊永遠一致，不會抄錯。

工作紙**全英**（唔中英對照），內容依 part 的 sections／questions 而定：
    · 抬頭：課題英文標題、Name／Class／Date／Score 欄、Total marks
    · 多項選擇題：英文題目＋A–D 選項＋獨立作答格（O 形格）
    · 快測題（question.quick）：
        input → 額外印一條填數空格；rows → 印紅／藍／剔表格（每行一個 ☐）
    · 長題目：英文題目＋各小問＋足夠書寫的橫線
    · 尾頁：Formula sheet and common traps（取自本章重點）
    · 頁尾：Page N
"""
import argparse
import json
import os
import re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

INK = (0x1F, 0x29, 0x37)
GREY = (0x6B, 0x72, 0x80)
BLUE = (0x1D, 0x4E, 0xD8)

ASCII_FONT = "Segoe UI"
EA_FONT = "Microsoft JhengHei"

# LaTeX → 純文字（工作紙唔跑 KaTeX，數式用 Unicode 就夠）
SUBS = {"0": "₀", "1": "₁", "2": "₂", "3": "₃", "4": "₄",
        "5": "₅", "6": "₆", "7": "₇", "8": "₈", "9": "₉",
        "n": "ₙ", "r": "ᵣ", "a": "ₐ", "k": "ₖ"}
SUP = {"0": "⁰", "1": "¹", "2": "²", "3": "³", "4": "⁴",
       "5": "⁵", "6": "⁶", "7": "⁷", "8": "⁸", "9": "⁹",
       "n": "ⁿ", "(": "⁽", ")": "⁾"}
REPL = {"\\,": " ", "\\;": " ", "\\!": "", "\\times": "×", "\\div": "÷",
        "\\cdot": "×", "\\le": "≤", "\\ge": "≥", "\\cdots": "…", "\\ldots": "…",
        "\\dots": "…", "\\lfloor": "⌊", "\\rfloor": "⌋", "\\pi": "π",
        "\\left": "", "\\right": "", "\\quad": "  ", "\\binom": "C",
        "\\text": ""}


def plain(s):
    """把 KaTeX 字串變成可以列印的純文字。"""
    if s is None:
        return ""
    t = str(s)
    for k, v in REPL.items():
        t = t.replace(k, v)
    t = t.replace("$", "").replace("{", "").replace("}", "")
    # 上標／下標
    res = []
    i = 0
    while i < len(t):
        ch = t[i]
        if ch in "^_" and i + 1 < len(t):
            kind = SUP if ch == "^" else SUBS
            group = t[i + 1]
            if group == "{":
                j = t.find("}", i + 2)
                if j > 0:
                    group, i = t[i + 2:j], j
            res.append("".join(kind.get(c, c) for c in group))
            i += 2
            continue
        res.append(ch)
        i += 1
    return "".join(res)


def en(node, key="en"):
    """只取英文（工作紙全英，唔中英對照）。"""
    if isinstance(node, dict):
        return node.get(key) or ""
    return node or ""


CJK = re.compile(r"[\u3000-\u303f\u3400-\u9fff\uff00-\uffef]+")


def ascii_only(s):
    """選項等欄位可能有「插空法 Slot-in」咁嘅中英混排，工作紙只要英文。"""
    t = CJK.sub(" ", s or "")
    return re.sub(r"\s{2,}", " ", t).strip()


def style_run(run, size=10.5, bold=False, color=INK, italic=False):
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor(*color)
    run.font.name = ASCII_FONT
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi"):
        rfonts.set(qn(attr), ASCII_FONT)
    rfonts.set(qn("w:eastAsia"), EA_FONT)
    return run


def para(doc, text="", size=10.5, bold=False, color=INK, space_after=4,
         space_before=0, align=None, indent=0, line=1.25, keep=False):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(space_before)
    pf.line_spacing = line
    pf.keep_with_next = keep          # 唔好將同一題拆到兩頁
    if indent:
        pf.left_indent = Cm(indent)
    if align is not None:
        p.alignment = align
    if text:
        style_run(p.add_run(text), size, bold, color)
    return p


def first_sentence(text):
    """工作紙只要引言第一句，唔想印五六行說明。"""
    if not text:
        return ""
    for sep in (". ", "? ", "! "):
        if sep in text:
            return text.split(sep)[0].strip() + sep.strip()
    return text


def cell_borders(cell, bottom=True, box=False):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement("w:" + edge)
        if (edge == "bottom" and bottom) or (box and edge in ("top", "left", "right")):
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), "4")
            el.set(qn("w:color"), "BFBFBF")
        else:
            el.set(qn("w:val"), "nil")
        borders.append(el)
    tc_pr.append(borders)


def keep_together(t):
    """整張表唔好拆到兩頁（逐行 cantSplit ＋ 逐格 keep_with_next）。"""
    for r in t.rows:
        tr_pr = r._tr.get_or_add_trPr()
        tr_pr.append(OxmlElement("w:cantSplit"))
        for c in r.cells:
            for p in c.paragraphs:
                p.paragraph_format.keep_with_next = True
    for c in t.rows[-1].cells:
        for p in c.paragraphs:
            p.paragraph_format.keep_with_next = False
    return t


def ruled_lines(doc, rows, row_cm=0.9):
    """書寫用橫線（只有下邊線）。"""
    t = doc.add_table(rows=rows, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    for r in t.rows:
        r.height = Cm(row_cm)
        cell_borders(r.cells[0], bottom=True)
        r.cells[0].width = Cm(17.0)
    return t


def option_box(cell, mark="○"):
    """作答用小格（單選 ○／多選 ☐）。"""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_run(p.add_run(mark), 11, False, GREY)


def header_block(doc, part, site_part):
    para(doc, en(part.get("title")), size=16, bold=True,
         space_after=2, align=WD_ALIGN_PARAGRAPH.CENTER)
    meta = site_part.get("meta") or {}
    bits = []
    if meta.get("paper"):
        bits.append(plain(meta["paper"]))
    if meta.get("marks"):
        bits.append("Total %s marks" % meta["marks"])
    if bits:
        para(doc, " · ".join(bits), size=10, color=GREY, space_after=8,
             align=WD_ALIGN_PARAGRAPH.CENTER)

    tb = doc.add_table(rows=1, cols=1)
    cell = tb.rows[0].cells[0]
    cell_borders(cell, bottom=True, box=True)
    cell.width = Cm(17.0)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    style_run(p.add_run("Name: ____________ Class: ________ "
                        "Date: ________ Score: ______ / %s"
                        % meta.get("marks", "")), 10, False, INK)

    intro = first_sentence(en(part.get("intro")).replace("{{math:0}}", ""))
    if intro:
        para(doc, intro, size=9.5, color=GREY, space_before=6, space_after=10)


def section_heading(doc, text_en, marks):
    p = para(doc, "", size=12.5, bold=True, space_before=8, space_after=4, keep=True)
    style_run(p.add_run(text_en), 12.5, True, BLUE)
    style_run(p.add_run(" (%s marks)" % marks), 10, False, GREY)


def add_input_line(doc, quick):
    """快測題的填數空格（Number of arrangements: ____________）。"""
    label = en(quick.get("input", {}).get("label")) or "Answer"
    tb = doc.add_table(rows=1, cols=1)
    cell = tb.rows[0].cells[0]
    cell_borders(cell, bottom=True, box=True)
    cell.width = Cm(17.0)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.left_indent = Cm(0.6)
    style_run(p.add_run("%s: " % label), 10.5, True, INK)
    style_run(p.add_run("____________________________"), 10.5, False, GREY)


def add_row_table(doc, quick):
    """快測題的剔選表格（head + rows + 每行一個 ☐）。"""
    head = [en(h) for h in quick.get("head") or []]
    rows = quick.get("rows") or []
    if not head or not rows:
        return
    value_cols = len(rows[0].get("cells") or [])
    # head 最後一欄如果係「剔／Tick」，就係打勾欄（否則每欄都對應一個 cells 值）
    has_tick = len(head) == value_cols + 1
    tick_col = 1 + value_cols if has_tick else None      # 0 欄係「(A)」咁嘅選項代號
    ncol = 1 + value_cols + (1 if has_tick else 0)
    t = doc.add_table(rows=1 + len(rows), cols=ncol)
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    for j in range(ncol):
        c = t.rows[0].cells[j]
        cell_borders(c, bottom=True, box=True)
        c.width = Cm(17.0 / ncol)
        pp = c.paragraphs[0]
        pp.paragraph_format.space_after = Pt(1)
        pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        style_run(pp.add_run(head[j - 1] if j else ""), 10.5, True, INK)
    for i, row in enumerate(rows, start=1):
        cells = t.rows[i].cells
        for j in range(ncol):
            cells[j].width = Cm(17.0 / ncol)
            pp = cells[j].paragraphs[0]
            pp.paragraph_format.space_after = Pt(1)
            pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if j == 0:
                cell_borders(cells[j], bottom=True, box=True)
                style_run(pp.add_run("(%s)" % row.get("opt", "")), 10.5, True, INK)
            elif j == tick_col:
                option_box(cells[j], mark="☐")
            else:
                cell_borders(cells[j], bottom=True, box=True)
                style_run(pp.add_run(plain(en(row["cells"][j - 1]))), 10.5)
        t.rows[i].height = Cm(0.7)
    keep_together(t)


def answer_sheet(doc, qs):
    """獨立作答格（DSE 卷同款：印喺卷頭；多選題用 ☐，單選用 ○）。"""
    sheet_qs = [q for q in qs if not (q.get("quick") or {}).get("input")]
    if not sheet_qs:
        return
    para(doc, "Answer sheet", size=10.5, bold=True, space_before=6,
         space_after=3, keep=True)
    cols = sorted({k for q in sheet_qs for k in q["options"].keys()})
    t = doc.add_table(rows=1 + len(sheet_qs), cols=1 + len(cols))
    heads = ["Q."] + cols
    for j, h in enumerate(heads):
        c = t.rows[0].cells[j]
        cell_borders(c, bottom=True, box=True)
        c.width = Cm(17.0 / len(cols))
        pp = c.paragraphs[0]
        pp.paragraph_format.space_after = Pt(0)
        pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        style_run(pp.add_run(h), 10, True, GREY)
    for i, q in enumerate(sheet_qs, start=1):
        row = t.rows[i]
        c = row.cells[0]
        cell_borders(c, bottom=True, box=True)
        pp = c.paragraphs[0]
        pp.paragraph_format.space_after = Pt(0)
        pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        style_run(pp.add_run(q["code"]), 10, True, INK)
        multi = bool((q.get("quick") or {}).get("multi"))
        for j in range(1, len(cols) + 1):
            cc = row.cells[j]
            cell_borders(cc, bottom=True, box=True)
            cc.width = Cm(17.0 / len(cols))
            option_box(cc, mark="☐" if multi else "○")
        for j in range(len(cols) + 1):
            row.cells[j].height = Cm(0.75)
    keep_together(t)
    para(doc, "", size=6, space_after=0)


def add_mc(doc, sec, qs):
    answer_sheet(doc, qs)
    title = en(sec.get("title")) or "Multiple choice"
    section_heading(doc, title, sec.get("marks") or sum(q["marks"] for q in qs))
    for q in qs:
        p = para(doc, "", space_after=1, space_before=6, keep=True)
        style_run(p.add_run("%s " % q["code"]), 11, True, INK)
        style_run(p.add_run("[%d mark%s]" % (q["marks"], "" if q["marks"] == 1 else "s")),
                  9.5, False, GREY)
        para(doc, plain(en(q["stem"])), size=10.5, space_after=2, indent=0.6, keep=True)
        quick = q.get("quick") or {}
        opts = q["options"]
        if not quick.get("rows"):
            # 有表格（紅／藍／剔）就唔好再印一次選項清單
            keys = sorted(opts.keys())
            t = doc.add_table(rows=(len(keys) + 1) // 2, cols=2)
            for i, k in enumerate(keys):
                c = t.rows[i // 2].cells[i % 2]
                cell_borders(c, bottom=False)
                c.width = Cm(8.5)
                pp = c.paragraphs[0]
                pp.paragraph_format.space_after = Pt(1)
                pp.paragraph_format.left_indent = Cm(0.6)
                style_run(pp.add_run("(%s) %s" % (k, ascii_only(plain(opts[k])))), 10.5)
        if quick.get("input"):
            add_input_line(doc, quick)
        if quick.get("rows"):
            add_row_table(doc, quick)
        para(doc, "", size=5, space_after=0)


def add_long(doc, sec, qs):
    total = sec.get("marks") or sum(q["marks"] for q in qs)
    title = en(sec.get("title")) or "Conventional questions"
    section_heading(doc, title, total)
    for q in qs:
        p = para(doc, "", space_after=1, space_before=8, keep=True)
        style_run(p.add_run("%s " % q["code"]), 11, True, INK)
        style_run(p.add_run("[%d marks]" % q["marks"]), 9.5, False, GREY)
        para(doc, plain(en(q["stem"])), size=10.5, space_after=2, indent=0.6, keep=True)
        for part in q.get("parts") or []:
            txt = " ".join(x for x in [part.get("label"), plain(en(part))] if x)
            sub = para(doc, "", space_after=1, indent=0.9, keep=True)
            style_run(sub.add_run(txt), 10.5)
            style_run(sub.add_run("  [%d mark%s]" % (part.get("marks", 0),
                                                     "" if part.get("marks", 0) == 1 else "s")),
                      9.5, False, GREY)
        rows = max(5, q["marks"] + 2)
        ruled_lines(doc, rows)
        para(doc, "", size=6, space_after=0)


def add_formula_page(doc, part):
    doc.add_page_break()
    para(doc, "Formula sheet and common traps",
         size=14, bold=True, color=BLUE, space_after=8)
    para(doc, "Permutations", size=11.5, bold=True, space_before=4, space_after=2)
    for s in ["Arrange all n objects: n!",
              "Arrange r objects chosen from n: P(n,r) = n! ÷ (n − r)!",
              "Tie-up (must sit together): arrange one unit fewer, then multiply by the orders inside the block",
              "Slot-in (must not sit together): seat the others first; gaps = number of people + 1, at most one person per gap"]:
        para(doc, "・" + s, size=10, space_after=2, indent=0.4)
    para(doc, "Combinations", size=11.5, bold=True, space_before=8, space_after=2)
    for s in ["Choose r objects from n (order ignored): C(n,r) = n! ÷ (r! × (n − r)!)",
              "P(n,r) = C(n,r) × r! (arrangements = choices × orders)",
              "Case splitting: add between cases, multiply within a case",
              "'At least ...': all ways − the ways that fail the condition (complement)",
              "'At most n': 0, 1, …, n gives n + 1 cases"]:
        para(doc, "・" + s, size=10, space_after=2, indent=0.4)
    para(doc, "Counting principles", size=11.5, bold=True, space_before=8, space_after=2)
    for s in ["Cases (choose one way): add; steps (do one after another): multiply",
              "'Or' = add between cases, but first check whether the cases overlap (if so, subtract the intersection)",
              "Passwords / codes: order matters, so use permutations; handle the restricted positions first"]:
        para(doc, "・" + s, size=10, space_after=2, indent=0.4)


def add_footer(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_run(p.add_run("Page "), 9, False, GREY)
    r = p.add_run()
    style_run(r, 9, False, GREY)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    r._element.addnext(fld)


def main():
    ap = argparse.ArgumentParser(description="由網站資料檔生成全英 Word 工作紙")
    ap.add_argument("--part", default="ch17-test", help="課題 id（預設 ch17-test）")
    ap.add_argument("--out", default=None, help="輸出 .docx 路徑")
    ap.add_argument("--no-formula", action="store_true", help="不印尾頁公式速查")
    args = ap.parse_args()

    src = os.path.join(ROOT, "data", "src", args.part + ".json")
    data = json.load(open(src, encoding="utf-8"))
    site = json.load(open(os.path.join(ROOT, "data", "src", "site.json"), encoding="utf-8"))
    site_part = next(p for p in site["site"]["parts"] if p["id"] == args.part)

    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = ASCII_FONT
    normal.font.size = Pt(10.5)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), EA_FONT)

    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.0)
    sec.top_margin = Cm(1.6)
    sec.bottom_margin = Cm(1.6)
    add_footer(sec)

    header_block(doc, data, site_part)
    for s in data.get("sections", []):
        qs = s.get("questions", [])
        if not qs:
            continue
        if qs[0]["type"] == "mc":
            add_mc(doc, s, qs)
        else:
            add_long(doc, s, qs)
    if not args.no_formula:
        add_formula_page(doc, data)

    out = args.out or os.path.join(ROOT, "docs", "worksheets",
                                   "%s-worksheet.docx" % args.part)
    os.makedirs(os.path.dirname(out), exist_ok=True)
    doc.save(out)
    print("已生成：%s（%d 題，全英）" % (
        out, sum(len(s.get("questions", [])) for s in data.get("sections", []))))


if __name__ == "__main__":
    main()